"""Escritor del libro registro LSI (app/fiscal_modelos/libros_lsi.py). La prueba de verdad es la del simulador de la AEAT (escenarios de test_fiscal_motor_303.py, grabados con
scripts/fiscal/grabar_oraculo.py): aquí se fija la forma del fichero que descubrimos que acepta."""
import io
from datetime import date as d, datetime

import openpyxl

from app.fiscal_modelos import libros_lsi as lsi
from app.fiscal_modelos.apuntes import Apunte, en_periodo, periodo_de
from tests.escenarios_303 import compra, venta


def _hojas(contenido: bytes):
    return openpyxl.load_workbook(io.BytesIO(contenido))


def test_tres_hojas_con_los_nombres_oficiales_y_datos_desde_la_fila_1_sin_cabeceras():
    wb = _hojas(lsi.escribir_xlsx([venta(10000), compra(5000)], ejercicio=2026, periodo="2T"))
    assert wb.sheetnames == ["EXPEDIDAS_INGRESOS", "RECIBIDAS_GASTOS", "BIENES-INVERSIÓN"]
    e, r = wb["EXPEDIDAS_INGRESOS"], wb["RECIBIDAS_GASTOS"]
    assert e.max_row == 1 and r.max_row == 1 and wb["BIENES-INVERSIÓN"].max_row == 1
    fila = [c.value for c in e[1]]
    assert fila[0] == 2026 and fila[1] == "2T" and fila[2:4] == ["B", "01"] and fila[5] == "F1" and fila[8] == datetime(2026, 5, 10)
    assert fila[20:24] == [121.0, 100.0, 21.0, 21.0]                      # total, base, tipo, cuota (euros con dos decimales)
    assert e.cell(1, 9).number_format == "dd/mm/yyyy"
    fila = [c.value for c in r[1]]
    assert fila[5] == "F1" and fila[12] == datetime(2026, 5, 12) and fila[26] == 50.0 and fila[28] == 10.5 and fila[29] == 10.5


def test_cada_fila_va_con_su_periodo_natural():
    ap = [venta(100, fecha=d(2026, 2, 1)), venta(100, numero="2", fecha=d(2026, 5, 1), fecha_operacion=d(2026, 3, 30)),
          compra(100, fecha=d(2026, 1, 2), fecha_recepcion=d(2026, 8, 1)), compra(100, numero="P2", fecha=d(2026, 5, 2), deducible_periodo_posterior=True, periodo_deduccion=(2026, "3T"))]
    wb = _hojas(lsi.escribir_xlsx(ap, ejercicio=2026, periodo="3T"))
    assert [c.value for c in wb["EXPEDIDAS_INGRESOS"]["B"]] == ["1T", "1T"]
    assert [c.value for c in wb["RECIBIDAS_GASTOS"]["B"]] == ["3T", "2T"]
    assert wb["RECIBIDAS_GASTOS"].cell(2, 23).value == "S" and (wb["RECIBIDAS_GASTOS"].cell(2, 24).value, wb["RECIBIDAS_GASTOS"].cell(2, 25).value) == (2026, "3T")


def test_sin_iva_repercutido_tipo_y_cuota_van_a_cero_no_vacios():
    wb = _hojas(lsi.escribir_xlsx([venta(5000, 0, 0, calificacion="S2")], ejercicio=2026, periodo="2T"))
    fila = [c.value for c in wb["EXPEDIDAS_INGRESOS"][1]]
    assert fila[18] == "S2" and fila[22] == 0 and fila[23] == 0 and fila[20] == 50.0          # con inversión del sujeto pasivo, total = base


def test_adquisicion_intracomunitaria_e_isp_llevan_total_igual_a_la_base():
    a = compra(20000, clave="09", contraparte_nif="FR40303265045", contraparte_nif_tipo="02", contraparte_pais="FR")
    b = compra(15000, isp=True)
    assert a.total_efectivo == 20000 and b.total_efectivo == 15000
    wb = _hojas(lsi.escribir_xlsx([a], ejercicio=2026, periodo="2T"))
    fila = [c.value for c in wb["RECIBIDAS_GASTOS"][1]]
    assert fila[15:18] == ["02", "FR", "FR40303265045"] and fila[25] == 200.0


def test_periodos():
    assert periodo_de(d(2026, 6, 30)) == "2T" and periodo_de(d(2026, 7, 1)) == "3T" and periodo_de(d(2026, 5, 3), mensual=True) == "05"
    assert en_periodo(d(2026, 6, 30), 2026, "2T") and not en_periodo(d(2025, 6, 30), 2026, "2T") and not en_periodo(None, 2026, "2T") and en_periodo(d(2026, 5, 3), 2026, "05")


def test_el_apunte_es_inmutable_y_con_reemplaza():
    a = venta(10000)
    b = a.con(base=500)
    assert a.base == 10000 and b.base == 500
    try:
        a.base = 1                                                          # type: ignore[misc]
    except Exception as e:                                                  # noqa: BLE001
        assert "frozen" in type(e).__name__.lower() or "cannot assign" in str(e)
    else:
        raise AssertionError("debería ser inmutable")
    assert isinstance(a, Apunte)


def test_ida_y_vuelta_de_todos_los_escenarios():
    """Lo que escribimos, lo leemos igual: cada escenario (los 23 que la AEAT aceptó) vuelve a ser el mismo conjunto de apuntes."""
    from tests.escenarios_303 import ESCENARIOS
    for nombre, (apuntes, ejercicio, periodo) in ESCENARIOS.items():
        leidos, ilegibles = lsi.leer_xlsx(lsi.escribir_xlsx(apuntes, ejercicio=ejercicio, periodo=periodo))
        assert ilegibles == [], nombre
        esperados = [a for a in apuntes if a.libro == "E"] + [a for a in apuntes if a.libro == "R"]          # el libro agrupa primero las emitidas
        assert len(leidos) == len(esperados), nombre
        for l, e in zip(leidos, esperados):
            assert (l.libro, l.tipo_factura, l.fecha_expedicion, l.base, l.cuota, l.clave, l.calificacion, l.exenta, l.contraparte_nif, l.isp, l.bien_inversion, l.periodo_deduccion) == \
                   (e.libro, e.tipo_factura, e.fecha_expedicion, e.base, e.cuota, e.clave, e.calificacion if e.calificacion == "S1" or e.calificacion else "", e.exenta, e.contraparte_nif, e.isp,
                    e.bien_inversion, e.periodo_deduccion), nombre


def test_las_cabeceras_de_las_plantillas_y_las_filas_vacias_se_saltan():
    import openpyxl
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("EXPEDIDAS_INGRESOS")
    ws.append(["LIBRO REGISTRO FACTURAS EXPEDIDAS"]); ws.append(["Ejercicio", "Periodo"]); ws.append([None] * 5)
    ws.append([2026, "2T", "B", "01", None, "F1", None, None, "10/05/2026", None, "A", "1", None, None, None, "B87654321", "CLIENTE SL", "01", "S1", None, "121,00", "100,00", "21,00", "21,00"])
    ws.append([2026, "2T", "B", "01", None, "F1", None, None, "no es fecha", None, "A", "2"])
    buf = io.BytesIO(); wb.save(buf)
    apuntes, ilegibles = lsi.leer_xlsx(buf.getvalue())
    assert [(a.numero, a.base, a.cuota, a.fecha_expedicion) for a in apuntes] == [("1", 10000, 2100, d(2026, 5, 10))]
    assert len(ilegibles) == 1 and ilegibles[0]["fila"] == 5 and "fecha no válida" in ilegibles[0]["motivo"]
