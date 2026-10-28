"""Criterio de caja (clave 07) en el 303: el IVA se devenga o se deduce al cobrar o pagar, en proporción, o el 31 de diciembre del año siguiente; casillas 62, 63, 74 y 75 de información."""
from datetime import date

import pytest

from app.fiscal_modelos import motor_303 as m
from tests.escenarios_303 import compra, venta

D = date


def _venta(cobros=(), **kw):
    return venta(100000, 2100, clave="07", cobros=tuple(cobros), fecha=kw.pop("fecha", D(2026, 5, 10)), **kw)       # 1.000 + 210 = 1.210 €


def test_sin_cobros_no_se_devenga_nada_pero_se_informa_de_lo_que_habria_devengado():
    r = m.calcular([_venta()], ejercicio=2026, periodo="2T", criterio_caja=True)
    assert r.calculable, r.bloqueos
    assert r.casillas.get("07", 0) == 0 and r.casillas.get("09", 0) == 0 and r.casillas["27"] == 0
    assert r.casillas["62"] == 100000 and r.casillas["63"] == 21000


def test_un_cobro_total_devenga_en_el_periodo_del_cobro():
    v = _venta([(D(2026, 8, 20), 121000)])
    assert m.calcular([v], ejercicio=2026, periodo="2T").casillas.get("09", 0) == 0
    r3 = m.calcular([v], ejercicio=2026, periodo="3T")
    assert r3.casillas["07"] == 100000 and r3.casillas["09"] == 21000 and r3.casillas["27"] == 21000 and "62" not in r3.casillas     # el 3T devenga pero la regla general no lo habría devengado ahí
    assert r3.origen["09"] == [0]


def test_cobros_parciales_reparten_base_y_cuota_en_proporcion_y_suman_exactamente():
    v = _venta([(D(2026, 5, 20), 30000), (D(2026, 8, 1), 50000), (D(2026, 11, 15), 41000)])
    partes = [m.calcular([v], ejercicio=2026, periodo=p) for p in ("2T", "3T", "4T")]
    assert [p.casillas.get("09", 0) for p in partes] == [5207, 8677, 7116] and sum(p.casillas.get("09", 0) for p in partes) == 21000
    assert sum(p.casillas.get("07", 0) for p in partes) == 100000


def test_lo_que_no_se_cobra_se_devenga_el_31_de_diciembre_del_ano_siguiente():
    v = _venta([(D(2026, 6, 1), 60500)])
    assert m.calcular([v], ejercicio=2026, periodo="4T").casillas.get("09", 0) == 0
    for periodo in ("1T", "2T", "3T"):
        r = m.calcular([v], ejercicio=2027, periodo=periodo)
        assert r.casillas.get("09", 0) == 0
    r = m.calcular([v], ejercicio=2027, periodo="4T")
    assert r.casillas["09"] == 10500 and r.casillas["07"] == 50000                           # la mitad que quedaba por cobrar
    assert m.calcular([v], ejercicio=2028, periodo="1T").casillas.get("09", 0) == 0


def test_una_devolucion_de_lo_cobrado_resta():
    v = _venta([(D(2026, 5, 20), 121000), (D(2026, 8, 5), -121000)])
    assert m.calcular([v], ejercicio=2026, periodo="2T").casillas["09"] == 21000
    r3 = m.calcular([v], ejercicio=2026, periodo="3T")
    assert r3.casillas["09"] == -21000 and r3.casillas["27"] == -21000


def test_cobros_de_mas_o_facturas_no_admitidas_bloquean():
    r = m.calcular([_venta([(D(2026, 5, 20), 130000)])], ejercicio=2026, periodo="2T")
    assert not r.calculable and any("más que el total" in b for b in r.bloqueos)
    for malo in (_venta(tipo_factura="R1"), _venta(calificacion="S2"), _venta(exenta="E1"), _venta(recargo_pct=520, recargo_cuota=5200)):
        assert not m.calcular([malo], ejercicio=2026, periodo="2T").calculable


def test_la_deduccion_de_una_recibida_con_caja_sigue_a_los_pagos():
    c = compra(50000, 2100, clave="07", cobros=((D(2026, 6, 20), 30000), (D(2026, 9, 10), 30500)), fecha=D(2026, 5, 12))      # 500 + 105 = 605 €
    r2 = m.calcular([c], ejercicio=2026, periodo="2T")
    assert r2.casillas["74"] == 50000 and r2.casillas["75"] == 10500                          # información: se habría deducido en el 2T
    assert r2.casillas["28"] == 24793
    assert r2.casillas["29"] == 5207
    r3 = m.calcular([c], ejercicio=2026, periodo="3T")
    assert r3.casillas["29"] == 5293 and "74" not in r3.casillas and r2.casillas["29"] + r3.casillas["29"] == 10500
    assert r2.casillas["45"] == 5207


def test_recibida_con_caja_inversion_prorrata_y_exclusiones():
    c = compra(50000, 2100, clave="07", cobros=((D(2026, 6, 20), 60500),), fecha=D(2026, 5, 12), bien_inversion=True)
    assert m.calcular([c], ejercicio=2026, periodo="2T").casillas["31"] == 10500
    r = m.calcular([compra(50000, 2100, clave="07", cobros=((D(2026, 6, 20), 60500),), fecha=D(2026, 5, 12))], ejercicio=2026, periodo="2T", prorrata_general=5000)
    assert r.casillas["29"] == 5250 and r.casillas["75"] == 10500                              # solo el 50 % de la cuota se deduce; la información lleva la cuota entera
    for malo in (compra(50000, 2100, clave="07", isp=True), compra(50000, 2100, clave="07", tipo_factura="R1"), compra(50000, 2100, clave="07", cobros=((D(2026, 6, 20), 70000),))):
        assert not m.calcular([malo], ejercicio=2026, periodo="2T").calculable


def test_el_titular_en_caja_con_facturas_sin_clave_07_recibe_un_aviso():
    r = m.calcular([venta(100000, 2100, fecha=D(2026, 5, 10))], ejercicio=2026, periodo="2T", criterio_caja=True)
    assert r.calculable and any("sin la clave 07" in a for a in r.avisos) and r.casillas["09"] == 21000


# --- validación, libro de la AEAT, persistencia, fichero y 390 -------------------------------------------------------------------------------------------
from app.fiscal_modelos import libros_lsi, motor_390, validacion_apuntes  # noqa: E402

HOY = D(2026, 10, 28)


def _errores(a):
    return [e.codigo for e in validacion_apuntes.errores(validacion_apuntes.validar_apunte(a, hoy=HOY))]


def test_validacion_del_criterio_de_caja():
    assert _errores(_venta([(D(2026, 5, 20), 121000)])) == []
    assert "cobros_exceden" in _errores(_venta([(D(2026, 5, 20), 130000)]))
    assert "cobro_negativo" in _errores(_venta([(D(2026, 5, 20), -100)]))
    assert "cobro_futuro" in _errores(_venta([(D(2026, 12, 20), 100)]))
    assert "caja_calificacion" in _errores(_venta(calificacion="N2")) and "caja_calificacion" in _errores(_venta(exenta="E5", calificacion="S1"))
    assert "caja_rectificativa" in _errores(_venta(tipo_factura="R1"))
    assert "caja_isp" in _errores(compra(50000, 2100, clave="07", isp=True))
    assert "cobros_sin_caja" in [e.codigo for e in validacion_apuntes.validar_apunte(venta(100000, 2100, cobros=((D(2026, 5, 20), 100),)), hoy=HOY)]


def test_el_libro_de_la_aeat_lleva_la_factura_una_vez_y_una_linea_sin_importes_por_cobro_y_se_lee_de_vuelta_en_una_sola_factura():
    e = _venta([(D(2026, 5, 20), 30000), (D(2026, 8, 1), 91000)])
    r = compra(50000, 2100, clave="07", cobros=((D(2026, 6, 20), 60500),), fecha=D(2026, 5, 12))
    contenido = libros_lsi.escribir_xlsx([e, r, venta(5000, 2100, numero="9")], ejercicio=2026, periodo="3T")
    import io
    import openpyxl
    hojas = openpyxl.load_workbook(io.BytesIO(contenido))
    ex, re_ = hojas["EXPEDIDAS_INGRESOS"], hojas["RECIBIDAS_GASTOS"]
    assert ex.max_row == 4 and re_.max_row == 2                                         # la factura una vez, una línea por cobro o pago y la otra factura
    assert ex.cell(1, 18).value == "07" and ex.cell(1, 22).value == 1000.0 and ex.cell(1, 23).value == 21.0 and ex.cell(1, 27).value is None
    for fila, importe in ((2, 300.0), (3, 910.0)):                                      # la línea del cobro identifica la factura y NO lleva total, base, tipo ni cuota (el simulador de la AEAT las rechaza)
        assert ex.cell(fila, 12).value == "1" and ex.cell(fila, 28).value == importe and ex.cell(fila, 29).value == "04"
        assert all(ex.cell(fila, c).value is None for c in (21, 22, 23, 24))
    assert ex.cell(2, 2).value == "2T" and ex.cell(3, 2).value == "3T"                  # cada línea, en el periodo del cobro
    assert re_.cell(2, 34).value == 605.0 and all(re_.cell(2, c).value is None for c in (26, 27, 28, 29))
    leidos, ilegibles = libros_lsi.leer_xlsx(contenido)
    assert not ilegibles and len(leidos) == 3
    caja = next(a for a in leidos if a.libro == "E" and a.clave == "07")
    assert caja.cobros == ((D(2026, 5, 20), 30000), (D(2026, 8, 1), 91000))
    assert next(a for a in leidos if a.libro == "R").cobros == ((D(2026, 6, 20), 60500),)


def test_un_cobro_sin_su_factura_en_el_libro_se_senala():
    contenido = libros_lsi.escribir_xlsx([_venta([(D(2026, 5, 20), 30000)])], ejercicio=2026, periodo="2T")
    import io
    import openpyxl
    libro = openpyxl.load_workbook(io.BytesIO(contenido))
    libro["EXPEDIDAS_INGRESOS"].delete_rows(1)                                       # se quita la factura y queda el cobro suelto
    salida = io.BytesIO(); libro.save(salida)
    leidos, ilegibles = libros_lsi.leer_xlsx(salida.getvalue())
    assert leidos == [] and ilegibles and "no está en el libro" in ilegibles[0]["motivo"]


def test_un_libro_de_otro_programa_que_repite_la_factura_en_cada_cobro_tambien_se_une():
    e = _venta([(D(2026, 5, 20), 30000), (D(2026, 8, 1), 91000)])
    repetida = libros_lsi.escribir_xlsx([e], ejercicio=2026, periodo="3T")
    import io
    import openpyxl
    libro = openpyxl.load_workbook(io.BytesIO(repetida))
    ws = libro["EXPEDIDAS_INGRESOS"]
    for fila, (f, i) in zip((2, 3), e.cobros):                                       # cada línea repite los importes de la factura, como lo escribíamos antes
        for c, v in ((21, 121000 / 100), (22, 1000.0), (23, 21.0), (24, 210.0)):
            ws.cell(fila, c).value = v
    salida = io.BytesIO(); libro.save(salida)
    leidos, ilegibles = libros_lsi.leer_xlsx(salida.getvalue())
    assert not ilegibles and len(leidos) == 1 and len(leidos[0].cobros) == 2


def _despacho(nombre="Despacho caja", **perfil):
    from app import db, fiscal_perfiles as fp
    from app.db_fiscal import crear_cliente_fiscal
    t = db.crear_tenant(nombre)
    clave = nombre.replace(" ", "").lower()
    u = db.crear_usuario_vinculado_a_kratos(f"{clave}@ejemplo.com", f"k-{clave}"); db.asignar_tenant(u, t); db.asignar_supervisor_tenant(u, True)
    c = crear_cliente_fiscal(t, "Empresa de prueba SL", nif="A12345674")
    fp.guardar_perfil(t, u, c, {"forma_juridica": "sociedad", **perfil}, vigente_desde="2026-01-01")
    return t, u, c


def test_de_los_cobros_al_fichero_303_con_sus_indicadores_y_casillas_de_informacion():
    import json
    from app import fiscal_nucleo as fn
    from app.fiscal_modelos import exportador_303 as ex
    t, u, c = _despacho(criterio_caja=1)
    a = fn.aprobar_apunte(t, u, c, _venta([(D(2026, 5, 20), 60500), (D(2026, 8, 1), 60500)]))
    fn.aprobar_apunte(t, u, c, compra(50000, 2100, clave="07", cobros=((D(2026, 6, 20), 60500),), fecha=D(2026, 5, 12)))
    fila = next(f for f in fn.listar_apuntes(t, u, c) if f["id"] == a["id"])
    assert fn.cobros_de_json(fila["cobros_json"]) == ((D(2026, 5, 20), 60500), (D(2026, 8, 1), 60500))
    dec = fn.calcular_declaracion(t, u, fn.crear_instantanea(t, u, c, ejercicio=2026, periodo="2T")["id"])
    cas = dec["casillas"]
    assert dec["bloqueos"] == [] and cas["07"] == 50000 and cas["09"] == 10500 and cas["62"] == 100000 and cas["63"] == 21000 and cas["74"] == 50000 and cas["75"] == 10500
    assert cas["28"] == 50000 and cas["29"] == 10500 and cas["71"] == 0
    fn.revisar_declaracion(t, u, dec["id"])
    fn.exportar_declaracion(t, u, dec["id"])
    datos, _ = fn.descargar_artefacto(t, u, dec["id"], "303")
    diseno = json.loads(ex.DISENO_303_2026.read_text())
    pagina1 = next(p for p in ex.descomponer(diseno, datos)["paginas"] if p["hoja"] == "DP30301")["campos"]
    assert pagina1[15] == "1" and pagina1[16] == "1"                                                        # acogido a la caja y destinatario de operaciones en caja
    leidas = ex.leer_casillas(diseno, datos)
    assert leidas["62"] == 100000 and leidas["63"] == 21000 and leidas["74"] == 50000 and leidas["75"] == 10500 and leidas["09"] == 10500
    # el 3T devenga lo cobrado en agosto
    dec3 = fn.calcular_declaracion(t, u, fn.crear_instantanea(t, u, c, ejercicio=2026, periodo="3T")["id"])
    assert dec3["casillas"]["09"] == 10500 and "62" not in dec3["casillas"]


def test_el_390_con_apuntes_de_caja_se_bloquea():
    r = motor_390.calcular([_venta([(D(2026, 5, 20), 121000)])], ejercicio=2026)
    assert not r.calculable and any("criterio de caja" in b for b in r.bloqueos)
