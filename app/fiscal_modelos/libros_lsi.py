"""Escritor del **libro registro de IVA** en el formato de importación de la AEAT (`.xlsx`, «LSI»): lo que acepta la opción «Libros de registro → Importar libro de registro» del modelo 303.

Formato descubierto contra el simulador público (docs/fiscal-vias-aeat.md): tres hojas con los nombres oficiales `EXPEDIDAS_INGRESOS`, `RECIBIDAS_GASTOS` y `BIENES-INVERSIÓN`; los datos empiezan en la **fila 1**
(SIN filas de cabecera: el simulador las leería como registros) y las columnas son las de las plantillas oficiales (`docs/fuentes-oficiales/aeat/LSI*.xlsx`). Fechas como fecha de Excel, importes con dos decimales.
Desde 2023 hay que importar el libro COMPLETO del ejercicio hasta el periodo (el 4T incluye 1T-3T)."""
from __future__ import annotations

import io
from datetime import datetime

from .apuntes import EXPEDIDAS, RECIBIDAS, Apunte, periodo_de

HOJAS = ("EXPEDIDAS_INGRESOS", "RECIBIDAS_GASTOS", "BIENES-INVERSIÓN")
ACTIVIDAD = ("B", "01")                      # código y tipo de actividad: B = actividad no incluida en A, C ni D; tipo 01 (la pareja que el simulador acepta sin pedir inmueble ni epígrafe IAE)


def _euros(centimos: int | None) -> float | None:
    return None if centimos is None else round(centimos / 100, 2)


def _pct(centesimas: int) -> float:
    return round(centesimas / 100, 2)


MEDIO_OTROS = "04"                      # «Medio utilizado» del cobro o pago: 01 transferencia, 02 cheque, 03 no se cobra, 04 otros medios, 05 domiciliación


def _fecha(d) -> datetime | None:
    return datetime(d.year, d.month, d.day) if d else None


def periodo_natural(a: Apunte, periodo_objetivo: str) -> str:
    """El periodo con el que se anota la fila en el libro: el de la fecha que manda en cada libro (operación o expedición en las expedidas; recepción o periodo de deducción en las recibidas)."""
    mensual = periodo_objetivo[-1] != "T"
    if a.libro == RECIBIDAS:
        return periodo_de(a.fecha_recepcion or a.fecha_expedicion, mensual)
    return periodo_de(a.fecha_operacion or a.fecha_expedicion, mensual)


def _entra(fecha, ejercicio: int, periodo: str) -> bool:
    """¿Un cobro o pago de esa fecha cabe en el libro del ejercicio hasta el periodo de la autoliquidación? Los posteriores se anotarán cuando ocurran."""
    mensual = periodo[-1] != "T"
    natural = periodo_de(fecha, mensual)
    return fecha.year == ejercicio and (int(natural) <= int(periodo) if mensual else int(natural[0]) <= int(periodo[0]))


# Columnas (base 0) que la línea de un cobro o pago NO lleva: total, base, tipo y cuota (y recargo y deducible). El simulador de la AEAT las rechaza («incompatible rellenar el campo tipo de IVA con los
# campos de cobro»; «si hay base imponible debe indicar la cuota»); la línea solo identifica la factura y aporta la fecha, el importe y el medio.
CAMPOS_DE_IMPORTE = {EXPEDIDAS: (20, 21, 22, 23, 24, 25), RECIBIDAS: (25, 26, 27, 28, 29)}


def _solo_cobro(f: dict[int, object], libro: str) -> dict[int, object]:
    return {c: v for c, v in f.items() if c not in CAMPOS_DE_IMPORTE[libro]}


def fila_expedida(a: Apunte, ejercicio: int, periodo: str, cobro=None) -> dict[int, object]:
    """Columnas (base 0) del libro de expedidas. `cobro`: `(fecha, importe)` de la línea de un cobro (criterio de caja)."""
    f = {0: ejercicio, 1: periodo, 2: ACTIVIDAD[0], 3: ACTIVIDAD[1], 5: a.tipo_factura, 8: _fecha(a.fecha_expedicion), 9: _fecha(a.fecha_operacion), 10: a.serie or None, 11: a.numero,
         15: a.contraparte_nif or None, 16: a.contraparte_nombre, 17: a.clave, 18: a.calificacion or None, 19: a.exenta or None, 20: _euros(a.total_efectivo), 21: _euros(a.base)}
    if a.contraparte_nif_tipo:
        f[13], f[14] = a.contraparte_nif_tipo, a.contraparte_pais or None
    f[22], f[23] = (_pct(a.tipo_iva), _euros(a.cuota)) if a.calificacion == "S1" else (0.0, 0.0)       # S2/N/exentas: tipo y cuota a cero, no vacíos
    if a.recargo_pct or a.recargo_cuota:
        f[24], f[25] = _pct(a.recargo_pct), _euros(a.recargo_cuota)
    if cobro:
        f[26], f[27], f[28] = _fecha(cobro[0]), _euros(cobro[1]), MEDIO_OTROS
    if a.inmueble_situacion:
        f[33], f[34] = a.inmueble_situacion, a.inmueble_referencia or None
    return f


def fila_recibida(a: Apunte, ejercicio: int, periodo: str, pago=None) -> dict[int, object]:
    """Columnas (base 0) del libro de recibidas. `pago`: `(fecha, importe)` de la línea de un pago (criterio de caja)."""
    f = {0: ejercicio, 1: periodo, 2: ACTIVIDAD[0], 3: ACTIVIDAD[1], 5: a.tipo_factura, 8: _fecha(a.fecha_expedicion), 9: _fecha(a.fecha_operacion), 10: f"{a.serie}{a.numero}",
         12: _fecha(a.fecha_recepcion or a.fecha_expedicion), 17: a.contraparte_nif or None, 18: a.contraparte_nombre, 19: a.clave, 20: "S" if a.bien_inversion else None,
         21: "S" if a.isp else None, 22: "S" if a.deducible_periodo_posterior else None, 25: _euros(a.total_efectivo), 26: _euros(a.base), 27: _pct(a.tipo_iva), 28: _euros(a.cuota),
         29: _euros(a.deducible)}
    if a.contraparte_nif_tipo:
        f[15], f[16] = a.contraparte_nif_tipo, a.contraparte_pais or None
    if a.periodo_deduccion:
        f[23], f[24] = a.periodo_deduccion
    if pago:
        f[32], f[33], f[34] = _fecha(pago[0]), _euros(pago[1]), MEDIO_OTROS
    if a.inmueble_situacion:
        f[39], f[40] = a.inmueble_situacion, a.inmueble_referencia or None
    return f


def escribir_xlsx(apuntes: list[Apunte], *, ejercicio: int, periodo: str) -> bytes:
    """El libro (expedidas y recibidas) como `.xlsx`. `periodo` es el de la autoliquidación objetivo («2T»); cada fila se anota con SU periodo natural (`periodo_natural`), así que el libro puede llevar todo el ejercicio."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, libro, fila in ((HOJAS[0], EXPEDIDAS, fila_expedida), (HOJAS[1], RECIBIDAS, fila_recibida), (HOJAS[2], None, None)):
        ws = wb.create_sheet(nombre)
        if libro is None:
            continue
        n = 0
        for a in apuntes:
            if a.libro != libro:
                continue
            filas = [fila(a, ejercicio, periodo_natural(a, periodo))]
            if a.clave == "07":                                                                     # criterio de caja: la factura va UNA vez; cada cobro o pago, en una línea propia del periodo en que se produce
                filas += [_solo_cobro(fila(a, ejercicio, periodo_de(c[0], periodo[-1] != "T"), c), libro) for c in sorted(a.cobros) if _entra(c[0], ejercicio, periodo)]
            for f in filas:
                n += 1
                for col, valor in f.items():
                    if valor is None or valor == "":
                        continue
                    celda = ws.cell(n, col + 1)
                    celda.value = valor
                    if isinstance(valor, datetime):
                        celda.number_format = "dd/mm/yyyy"
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


# --- Lectura: de un libro LSI a apuntes --------------------------------------------------------------------------------------------------------------
def _f(v):
    """Fecha de una celda (fecha de Excel o «dd/mm/aaaa» / «aaaa-mm-dd»); `None` si no hay."""
    from datetime import date
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    texto = str(v).strip()
    try:
        if "/" in texto:
            d, m, a = texto.split("/")
            return date(int(a), int(m), int(d))
        return date.fromisoformat(texto[:10])
    except ValueError:
        raise ValueError(f"fecha no válida: {texto!r}") from None


def _cobro(fecha, importe, medio) -> tuple:
    """El cobro o pago de una línea del libro. El medio «03» (no se cobra) marca el límite de devengo del 31 de diciembre: no es un cobro."""
    if fecha in (None, "") or importe in (None, "") or _s(medio) == "03":
        return ()
    return ((_f(fecha), _c(importe)),)


def _c(v) -> int:
    """Importe de una celda en céntimos (número o texto con coma)."""
    from decimal import ROUND_HALF_UP, Decimal
    if v in (None, ""):
        return 0
    if isinstance(v, (int, float)):
        return int((Decimal(str(v)) * 100).quantize(Decimal(1), rounding=ROUND_HALF_UP))
    t = str(v).strip().replace(" ", "")
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    return int((Decimal(t) * 100).quantize(Decimal(1), rounding=ROUND_HALF_UP))


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def leer_xlsx(contenido: bytes) -> tuple[list[Apunte], list[dict]]:
    """Los apuntes de un libro LSI (el que escribimos nosotros o el de otro programa en el formato oficial) y las filas que NO se han podido leer (`{hoja, fila, motivo}`). Se saltan las filas de cabecera
    (las de las plantillas oficiales) y las vacías. No valida las reglas de la AEAT: eso lo hace `validacion_apuntes`."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    apuntes: list[Apunte] = []
    origen: list[tuple[str, int]] = []                                              # hoja y fila de cada apunte, para señalar un cobro sin factura
    ilegibles: list[dict] = []
    for nombre, libro in ((HOJAS[0], EXPEDIDAS), (HOJAS[1], RECIBIDAS)):
        if nombre not in wb.sheetnames:
            continue
        for n, fila in enumerate(wb[nombre].iter_rows(values_only=True), 1):
            c = list(fila) + [None] * 45
            if not any(x not in (None, "") for x in c[:30]):
                continue
            try:
                if not isinstance(c[0], (int, float)) and not _s(c[0]).isdigit():
                    continue                                                    # cabecera o texto explicativo de la plantilla
                origen.append((nombre, n))
                if libro == EXPEDIDAS:
                    apuntes.append(Apunte(EXPEDIDAS, _s(c[5]), _f(c[8]), numero=_s(c[11]), serie=_s(c[10]), clave=_s(c[17]) or "01", base=_c(c[21]), tipo_iva=_c(c[22]), cuota=_c(c[23]),
                                          total=_c(c[20]) if c[20] not in (None, "") else None, fecha_operacion=_f(c[9]), contraparte_nombre=_s(c[16]), contraparte_nif=_s(c[15]),
                                          contraparte_nif_tipo=_s(c[13]), contraparte_pais=_s(c[14]), calificacion=_s(c[18]), exenta=_s(c[19]), inmueble_situacion=_s(c[33]), inmueble_referencia=_s(c[34]),
                                          recargo_pct=_c(c[24]) if c[24] not in (None, "") else 0, recargo_cuota=_c(c[25]) if c[25] not in (None, "") else 0, cobros=_cobro(c[26], c[27], c[28])))
                else:
                    ded = (int(float(c[23])), _s(c[24])) if c[23] not in (None, "") and c[24] not in (None, "") else None
                    apuntes.append(Apunte(RECIBIDAS, _s(c[5]), _f(c[8]), numero=_s(c[10]), clave=_s(c[19]) or "01", base=_c(c[26]), tipo_iva=_c(c[27]), cuota=_c(c[28]),
                                          total=_c(c[25]) if c[25] not in (None, "") else None, fecha_operacion=_f(c[9]), fecha_recepcion=_f(c[12]), contraparte_nombre=_s(c[18]), contraparte_nif=_s(c[17]),
                                          contraparte_nif_tipo=_s(c[15]), contraparte_pais=_s(c[16]), bien_inversion=_s(c[20]).upper() == "S", isp=_s(c[21]).upper() == "S",
                                          cuota_deducible=_c(c[29]) if c[29] not in (None, "") else None, deducible_periodo_posterior=_s(c[22]).upper() == "S", periodo_deduccion=ded,
                                          inmueble_situacion=_s(c[39]), inmueble_referencia=_s(c[40]), cobros=_cobro(c[32], c[33], c[34])))
            except (ValueError, TypeError, ArithmeticError) as e:
                origen = origen[:len(apuntes)]
                ilegibles.append({"hoja": nombre, "fila": n, "motivo": str(e)[:150]})
    unidos, sueltos = _unir_cobros(apuntes)
    ilegibles += [{"hoja": origen[i][0], "fila": origen[i][1], "motivo": "cobro o pago de una factura que no está en el libro"} for i in sueltos]
    return unidos, ilegibles


def _unir_cobros(apuntes: list[Apunte]) -> tuple[list[Apunte], list[int]]:
    """En el libro del criterio de caja la factura va una vez y cada cobro o pago, en una línea propia SIN importes (formato de la AEAT); un libro de otro programa puede repetir además todos los datos de la
    factura en cada línea. Se vuelve a una sola factura con todos sus cobros. Devuelve también la posición (en `apuntes`) de las líneas de cobro cuya factura no está en el libro."""
    salida: list[Apunte] = []
    vistos: dict[tuple, int] = {}
    cobros_aparte: list[int] = []
    for i, a in enumerate(apuntes):
        if a.clave != "07":
            salida.append(a)
        elif a.cobros and not a.base and not a.cuota and not a.total:
            cobros_aparte.append(i)                                          # va después, esté donde esté su factura
        else:
            clave = (a.libro, a.tipo_factura, a.fecha_expedicion, a.serie, a.numero, a.contraparte_nif)
            if clave in vistos:
                previo = salida[vistos[clave]]
                salida[vistos[clave]] = previo.con(cobros=tuple(sorted(set(previo.cobros) | set(a.cobros))))
            else:
                vistos[clave] = len(salida)
                salida.append(a)
    sueltos: list[int] = []
    for i in cobros_aparte:
        a = apuntes[i]
        clave = (a.libro, a.tipo_factura, a.fecha_expedicion, a.serie, a.numero, a.contraparte_nif)
        if clave in vistos:
            previo = salida[vistos[clave]]
            salida[vistos[clave]] = previo.con(cobros=tuple(sorted(set(previo.cobros) | set(a.cobros))))
        else:
            sueltos.append(i)
    return salida, sueltos
