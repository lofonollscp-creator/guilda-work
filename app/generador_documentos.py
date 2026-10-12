"""Genera documentos (CSV/XLSX/DOCX/PDF) a partir de una tabla de datos ya
construida -- agnóstico de origen: el Asistente IA reúne los datos con las
tools de lectura que ya tiene (tareas/correo/fiscal/facturas/CRM/hojas...)
y le pasa el resultado ya en forma de columnas+filas a la tool
`generar_documento_al_chat` (app/mcp_tools.py), que llama a las funciones
de aquí. Mismo criterio "puro Python, sin binario de sistema externo" que
app/fichaje_export.py:a_pdf (reportlab) -- ver requirements.txt."""
import csv
import io

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import db, informes

# Tope de filas por documento -- evita que una tabla descontrolada genere
# un archivo enorme o cuelgue el proceso generándolo. Si el LLM pide más,
# se trunca aquí y se avisa en `notas` (nunca falla en silencio).
FILAS_MAXIMAS = 2000


def _truncar(filas: list[list], notas: str | None) -> tuple[list[list], str | None]:
    if len(filas) <= FILAS_MAXIMAS:
        return filas, notas
    aviso = f"(truncado a las primeras {FILAS_MAXIMAS} de {len(filas)} filas)"
    return filas[:FILAS_MAXIMAS], f"{notas} {aviso}" if notas else aviso


def tabla_a_csv(titulo: str, columnas: list[str], filas: list[list], notas: str | None = None) -> bytes:
    filas, notas = _truncar(filas, notas)
    buf = io.StringIO()
    buf.write(f"# {titulo}\n")
    buf.write(f"# Generado: {db.now_iso()}\n")
    if notas:
        buf.write(f"# {notas}\n")
    writer = csv.writer(buf)
    writer.writerow(columnas)
    for fila in filas:
        writer.writerow(fila)
    # BOM UTF-8 para que Excel en Windows abra los acentos bien de entrada
    # en vez de tener que elegir la codificación a mano.
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


def tabla_a_xlsx(titulo: str, columnas: list[str], filas: list[list], notas: str | None = None) -> bytes:
    filas, notas = _truncar(filas, notas)
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31] or "Datos"  # Excel no admite nombres de hoja de más de 31 caracteres

    ws.append(columnas)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        ws.append(fila)
    ws.freeze_panes = "A2"

    anchos = [len(str(c)) for c in columnas]
    for fila in filas:
        for i, valor in enumerate(fila):
            anchos[i] = max(anchos[i], len(str(valor)) if valor is not None else 0)
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(ancho + 2, 60)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def tabla_a_docx(titulo: str, columnas: list[str], filas: list[list], notas: str | None = None) -> bytes:
    filas, notas = _truncar(filas, notas)
    doc = Document()
    doc.add_heading(titulo, level=1)
    if notas:
        doc.add_paragraph(notas)

    tabla = doc.add_table(rows=1, cols=len(columnas))
    tabla.style = "Light Grid Accent 1"
    for celda, nombre in zip(tabla.rows[0].cells, columnas):
        celda.text = str(nombre)
        for parrafo in celda.paragraphs:
            for run in parrafo.runs:
                run.bold = True
    for fila in filas:
        celdas = tabla.add_row().cells
        for celda, valor in zip(celdas, fila):
            celda.text = "" if valor is None else str(valor)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def tabla_a_pdf(titulo: str, columnas: list[str], filas: list[list], notas: str | None = None, despacho: str | None = None) -> bytes:
    """PDF de una tabla con la capa común de informes (app/informes.py): cabecera, pie con paginación, fuente con soporte de catalán y francés."""
    filas, notas = _truncar(filas, notas)
    return informes.construir_pdf(titulo, [("tabla", columnas, filas)], despacho=despacho, subtitulo=notas, apaisado=len(columnas) > 6)


def texto_a_pdf(titulo: str, texto: str, notas: str | None = None, despacho: str | None = None) -> bytes:
    """PDF de un texto libre (p. ej. un informe redactado por la IA): un párrafo por línea, con todo escapado."""
    return informes.construir_pdf(titulo, [("parrafo", texto)], despacho=despacho, subtitulo=notas)
